"""DOCX loader for SharePoint procedures, work instructions and records.

- The document control table (Document no. / Revision / Effective date / Next review / Owner, or the
  Kyltek 'Nummer / No.' variant) becomes metadata.
- Headings: 'Heading N' styles, or numbered short paragraphs ('5.1 Eerste beoordeling ...').
- Numbered step paragraphs ('5.6.5<TAB>[F] If no leak ...') and procedure step tables
  (Step | Action | Verify) become one block per step, so a citation can point at '§6.2 step 7'.
- One-cell tables are WARNING / NOTE / STOP boxes.
"""

import re
import zipfile
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.types import Block, ParsedDocument, Section

CONTROL_KEYS = {
    "document no.": "doc_number",
    "nummer / no.": "doc_number",
    "revision": "revision",
    "versie / version": "revision",
    "effective date": "effective_date",
    "datum / date": "effective_date",
    "next review": "next_review",
    "owner": "owner",
    "approved by": "approved_by",
    "goedgekeurd door / approved by": "approved_by",
    "opgesteld door / prepared by": "prepared_by",
    "status": "status",
}
NUMBERED_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,2})\.?\s+(\S.{0,90})$")
NUMBERED_STEP_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){1,3})\t\s*(.+)$", re.S)
SIGNAL_RE = re.compile(r"^(DANGER|WARNING|CAUTION|STOP|NOTE|IMPORTANT|LET OP)", re.I)


def _clean(text: str) -> str:
    return re.sub(r"[  ]+", " ", text).strip()


def _control_table(table: Table) -> dict | None:
    cells = [[_clean(c.text) for c in row.cells] for row in table.rows]
    meta = {}
    for row in cells:
        for i in range(0, len(row) - 1, 2):
            key = CONTROL_KEYS.get(row[i].lower())
            if key and row[i + 1]:
                meta[key] = row[i + 1]
    return meta if "doc_number" in meta else None


def load_docx(path: Path) -> ParsedDocument:
    out = ParsedDocument()
    if path.stat().st_size == 0:
        out.parse_status = "empty"
        out.parse_notes.append("file is empty (0 bytes)")
        return out
    if not zipfile.is_zipfile(path):
        out.parse_status = "failed"
        out.parse_notes.append("not a valid .docx (zip) file")
        return out

    document = docx.Document(str(path))
    sec0 = document.sections[0]
    header = " ".join(p.text for p in sec0.header.paragraphs if p.text.strip())
    footer = " ".join(p.text for p in sec0.footer.paragraphs if p.text.strip())
    out.metadata = {"header": header, "footer": footer}

    sections: list[Section] = [Section(heading_path=["(start)"])]
    stack: list[tuple[int, str]] = []

    def current() -> Section:
        return sections[-1]

    def open_section(title: str, number: str | None, depth: int) -> None:
        nonlocal stack
        stack = [s for s in stack if s[0] < depth] + [(depth, title)]
        sections.append(Section(heading_path=[t for _, t in stack], section_no=number))

    def add(kind: str, text: str, extra_loc: dict | None = None, merge: bool | str = True) -> None:
        cur = current()
        loc = {"section": cur.section_no, "heading": cur.heading_path[-1]}
        if extra_loc:
            loc.update(extra_loc)
        cur.blocks.append(
            Block(kind=kind, text=text, locator=loc,
                  merge_key=merge if isinstance(merge, str) else (f"sec:{len(sections)}" if merge else None))
        )

    table_idx = 0
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(element, document)
            text = _clean(para.text) if "\t" not in para.text else para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            m_heading = NUMBERED_HEADING_RE.match(text)
            if style.startswith("Heading"):
                level = int(re.sub(r"\D", "", style) or 1)
                number = m_heading.group(1) if m_heading else None
                open_section(text, number, number.count(".") + 1 if number else level)
                continue
            m_step = NUMBERED_STEP_RE.match(text)
            if m_step:
                add("step", f"{m_step.group(1)} {_clean(m_step.group(2))}", {"step": m_step.group(1)}, merge=False)
                continue
            if m_heading and len(text) <= 90 and not text.endswith(".") and len(m_heading.group(2).split()) <= 10:
                number = m_heading.group(1)
                open_section(text, number, number.count(".") + 1)
                continue
            kind = "warning" if SIGNAL_RE.match(text) else "prose"
            prefix = "• " if "List" in style else ""
            add(kind, prefix + _clean(text))
        elif tag == "tbl":
            table = Table(element, document)
            meta = _control_table(table)
            if meta:
                out.metadata.update(meta)
                continue
            rows = [[_clean(c.text) for c in row.cells] for row in table.rows]
            # merged cells repeat their text in python-docx; collapse consecutive duplicates
            rows = [[c for i, c in enumerate(r) if i == 0 or c != r[i - 1]] for r in rows]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            if len(rows) == 1 and len(rows[0]) == 1:
                text = rows[0][0]
                add("warning" if SIGNAL_RE.match(text) else "prose", text, merge=False)
                continue
            header, body = rows[0], rows[1:]
            is_steps = header and header[0].lower() in {"step", "stap"}
            for r_i, row in enumerate(body, start=1):
                if not any(row):
                    continue
                text = " | ".join(f"{h}: {v}" for h, v in zip(header, row) if v)
                if is_steps:
                    add("step", f"Step {text.split('|')[0].split(':', 1)[1].strip()}: " + " | ".join(
                        f"{h}: {v}" for h, v in zip(header[1:], row[1:]) if v),
                        {"step": row[0], "table": table_idx}, merge=False)
                else:
                    add("table_row", text, {"table": table_idx, "row": r_i, "label": row[0]},
                        merge=False if len(header) == 2 else f"tbl:{len(sections)}:{table_idx}")
            table_idx += 1

    out.sections = [s for s in sections if s.blocks]
    if not out.sections:
        out.parse_status = "failed"
        out.parse_notes.append("no text extracted")
    return out
